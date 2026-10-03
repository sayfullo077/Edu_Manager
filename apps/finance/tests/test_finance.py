from datetime import date
from decimal import Decimal

import pytest
from django.core.exceptions import ValidationError
from django.urls import reverse
from django.utils import timezone

from apps.accounts.models import Role
from apps.common.templatetags.ui import money
from apps.contracts.models import Contract
from apps.finance import selectors
from apps.finance.models import Account, ExpenseCategory, Invoice, Payment, Transaction
from apps.finance.services import cash, expenses, invoices, ledger, payments
from apps.people.models import Guardian, Student, StudentGuardian
from conftest import make_staff

D = Decimal


# ---------- Proratsiya (maktab Excel'idagi haqiqiy misollar) ----------

@pytest.mark.parametrize("fee, start, expected", [
    ("1750000", date(2026, 9, 2), "1750000"),   # 1–5-kunlar: to'liq oy
    ("1750000", date(2026, 9, 15), "933000"),   # 16/30 → 933 333 → 933 000
    ("1575000", date(2026, 9, 17), "735000"),   # 14/30, 10% chegirmali
    ("1750000", date(2026, 9, 22), "525000"),   # 9/30
    ("1850000", date(2026, 9, 17), "863000"),   # rus sinfi, 14/30
])
def test_first_month_proration_matches_school_excel(fee, start, expected):
    assert invoices.first_month_amount(D(fee), start) == D(expected)


# ---------- Fixture'lar ----------

@pytest.fixture
def student(branch, school_class, year):
    s = Student.objects.create(branch=branch, last_name="Aliyev", first_name="Vali", birth_date=date(2016, 5, 1),
                               gender="M", grade=4, school_class=school_class, joined_at=date(2026, 9, 15))
    g = Guardian.objects.create(last_name="Aliyev", first_name="Anvar", phone="998901112233")
    StudentGuardian.objects.create(student=s, guardian=g, relation="father", is_primary=True)
    return s


@pytest.fixture
def contract(student, year, reception):
    c = Contract.objects.create(
        branch=student.branch, academic_year=year, student=student, guardian=student.guardians.first(),
        full_tariff=D("1750000"), start_date=date(2026, 9, 15), end_date=date(2027, 6, 30),
        status=Contract.Status.SIGNED, created_by=reception)
    invoices.generate_for_contract(c)
    return c


def pay(student, amount, by, method=Payment.Method.CARD, **kw):
    kw.setdefault("card_network", "humo" if method == Payment.Method.CARD else "")
    return payments.accept_payment(student=student, amount=D(amount), method=method, by=by, **kw)


# ---------- To'lov grafigi ----------

def test_contract_generates_ten_months_with_prorated_first(contract):
    inv = list(contract.invoices.order_by("month"))
    assert len(inv) == 10
    assert inv[0].amount == D("933000") and inv[0].waived == D("817000")
    assert all(i.amount == D("1750000") for i in inv[1:])
    assert inv[0].due_date == date(2026, 9, 25)  # 15-sentabrda kelgan: +10 kun muhlat
    assert inv[1].due_date == date(2026, 10, 10)
    invoices.generate_for_contract(contract)  # idempotent
    assert contract.invoices.count() == 10


def test_signing_via_service_generates_invoices_through_current_month(student, year, reception, monkeypatch):
    monkeypatch.setattr(timezone, "localdate", lambda *a: date(2026, 11, 20))
    from apps.contracts.services import contracts as contract_service
    c = Contract.objects.create(branch=student.branch, academic_year=year, student=student,
                                guardian=student.guardians.first(), full_tariff=D("1750000"),
                                start_date=date(2026, 9, 15), end_date=date(2027, 6, 30),
                                status=Contract.Status.SENT, created_by=reception)
    from apps.accounts.services import otp
    otp.issue_code("998901112233", "contract", subject=c.otp_subject,
                   extra={"student": "x", "number": c.number, "fee": "1"})
    from apps.accounts.infrastructure.sms import LocmemSMSBackend
    code = LocmemSMSBackend.outbox[-1][1].rsplit(" ", 1)[-1]
    contract_service.confirm_with_code(c, code=code, by=reception, ip="1.1.1.1")
    # Aralash usul: imzolanganda faqat joriy oygacha (sen, okt, noy), qolgani oyma-oy / "Grafik" orqali
    assert [f"{i.month:%m}" for i in c.invoices.order_by("month")] == ["09", "10", "11"]


# ---------- To'lovlar ----------

def test_payment_allocates_fifo_across_months(contract, student, reception):
    pay(student, "2000000", reception)
    first, second = contract.invoices.order_by("month")[:2]
    assert first.status == Invoice.Status.PAID and first.paid == D("933000")
    assert second.status == Invoice.Status.PARTIAL and second.paid == D("1067000")


def test_overpayment_rejected(contract, student, reception):
    debt = payments.outstanding(student)
    with pytest.raises(ValidationError):
        pay(student, debt + 1, reception)
    pay(student, debt, reception)  # to'liq yil oldindan — mumkin
    assert payments.outstanding(student) == 0


def test_cash_requires_open_session(contract, student, reception):
    with pytest.raises(ValidationError):
        pay(student, "100000", reception, method=Payment.Method.CASH)
    cash.open_session(student.branch, by=reception)
    p = pay(student, "100000", reception, method=Payment.Method.CASH)
    assert p.cash_session is not None


def test_card_commission_and_balances(contract, student, reception):
    pay(student, "1000000", reception)
    terminal = ledger.get_account(student.branch, Account.Kind.TERMINAL)
    t = Transaction.objects.get(account=terminal)
    assert t.commission == D("2000")  # 0.2%
    assert ledger.balance(terminal) == D("998000")


def test_reverse_restores_invoices_and_balance(contract, student, reception):
    p = pay(student, "933000", reception, method=Payment.Method.TRANSFER)
    payments.reverse_payment(p, reason="Xato kiritildi", by=reception)
    first = contract.invoices.order_by("month").first()
    assert first.paid == 0 and first.status == Invoice.Status.PENDING
    assert ledger.balance(ledger.get_account(student.branch, Account.Kind.BANK)) == 0
    assert Payment.objects.get().status == Payment.Status.REVERSED
    with pytest.raises(ValidationError):
        payments.reverse_payment(p, reason="yana", by=reception)
    assert Payment.objects.count() == 1  # hech narsa o'chirilmadi


def test_cancel_contract_cancels_future_unpaid(contract, student, reception):
    from apps.contracts.services import contracts as contract_service
    pay(student, "933000", reception)
    contract_service.cancel_contract(contract, reason="Ketdi", by=reception)
    statuses = list(contract.invoices.order_by("month").values_list("status", flat=True))
    assert statuses[0] == Invoice.Status.PAID
    today = timezone.localdate()
    future = contract.invoices.filter(month__gt=today)
    assert future.exists() and all(i.status == Invoice.Status.CANCELLED for i in future)


# ---------- Kassa va xarajatlar ----------

@pytest.fixture
def category(db):
    return ExpenseCategory.objects.create(name="Oziq-ovqat", kind="fixed")


def test_cash_session_close_with_shortage_records_adjustment(contract, student, reception):
    session = cash.open_session(student.branch, by=reception)
    pay(student, "500000", reception, method=Payment.Method.CASH)
    closed = cash.close_session(session, counted=D("490000"), by=reception)
    assert closed.expected_cash == D("500000") and closed.difference == D("-10000")
    assert Transaction.objects.filter(kind=Transaction.Kind.ADJUSTMENT, direction="out", amount=D("10000")).exists()
    assert ledger.balance(ledger.get_account(student.branch, Account.Kind.CASH)) == D("490000")


def test_only_one_open_session(branch, reception):
    cash.open_session(branch, by=reception)
    with pytest.raises(ValidationError):
        cash.open_session(branch, by=reception)


def test_expense_rules(branch, reception, category, contract, student):
    today = timezone.localdate()
    with pytest.raises(ValidationError):  # kassa yopiq
        expenses.record_expense(branch=branch, category=category, amount=D("1000"), account_kind="cash",
                                spent_at=today, description="x", by=reception)
    cash.open_session(branch, by=reception)
    with pytest.raises(ValidationError):  # kassada pul yo'q
        expenses.record_expense(branch=branch, category=category, amount=D("1000"), account_kind="cash",
                                spent_at=today, description="x", by=reception)
    pay(student, "300000", reception, method=Payment.Method.CASH)
    expenses.record_expense(branch=branch, category=category, amount=D("100000"), account_kind="cash",
                            spent_at=today, description="Non", by=reception)
    expenses.set_limit(branch=branch, category=category, month=today, limit=D("400000"), by=reception)
    b = selectors.budget(branch, today)
    row = b["fixed"][0]
    assert (row.limit, row.spent, row.remaining, row.percent) == (D("400000"), D("100000"), D("300000"), 25)


def test_month_summary(contract, student, reception):
    month = date(2026, 9, 1)
    pay(student, "500000", reception, paid_at=timezone.make_aware(timezone.datetime(2026, 9, 20, 10)))
    s = selectors.month_summary(student.branch, month)
    assert s.expected == D("933000") and s.collected_for_month == D("500000") and s.received == D("500000")
    assert s.paid_percent == 54


# ---------- Sahifalar va ruxsatlar ----------

def test_reception_home_is_dashboard(staff_client, contract):
    resp = staff_client.get(reverse("core:home"))
    assert resp.status_code == 200 and "Filial byudjeti" in resp.content.decode()


def test_head_teacher_has_no_finance_access(client, branch, contract, student):
    client.force_login(make_staff(branch, Role.HEAD_TEACHER, "998900000103"))
    for url in (reverse("finance:dashboard"), reverse("finance:cashbox"),
                reverse("finance:payment_create", args=[student.pk])):
        assert client.get(url).status_code == 403
    page = client.get(reverse("people:student_detail", args=[student.pk])).content.decode()
    assert "To'lov qabul qilish" not in page


def test_payment_via_page(staff_client, contract, student):
    resp = staff_client.post(reverse("finance:payment_create", args=[student.pk]),
                             {"amount": "933000", "method": "transfer"})
    assert resp.status_code == 302
    assert contract.invoices.order_by("month").first().status == Invoice.Status.PAID


# ---------- Yon panel (xarajat) ----------

def test_expense_drawer_error_rerenders_dashboard_with_drawer(staff_client, category):
    resp = staff_client.post(reverse("finance:expense_create"),
                             {"from": "drawer", "category": category.pk, "amount": "", "account_kind": "bank",
                              "description": "x", "spent_at": timezone.localdate().isoformat()})
    page = resp.content.decode()
    assert resp.status_code == 400
    assert 'class="drawer open"' in page and "Filial byudjeti" in page
    assert page.count('id="expense-drawer"') == 1


def test_expense_drawer_success_redirects_to_dashboard(staff_client, category):
    resp = staff_client.post(reverse("finance:expense_create"),
                             {"from": "drawer", "category": category.pk, "amount": "150000", "account_kind": "bank",
                              "description": "Internet", "spent_at": timezone.localdate().isoformat()})
    assert resp.status_code == 302 and resp.url.startswith(reverse("finance:dashboard"))


def test_dashboard_has_no_header_buttons_and_has_chart_hooks(staff_client, contract):
    page = staff_client.get(reverse("finance:dashboard")).content.decode()
    assert "data-line-chart" in page and "data-donut" in page and "hover-col" in page
    assert 'href="/cashbox/" class="btn' not in page



# ---------- Grafik oyma-oy: "Grafik" oynasi va cron ----------

@pytest.fixture
def partial_contract(student, year, reception):
    c = Contract.objects.create(
        branch=student.branch, academic_year=year, student=student, guardian=student.guardians.first(),
        full_tariff=D("1750000"), start_date=date(2026, 9, 15), end_date=date(2027, 6, 30),
        status=Contract.Status.SIGNED, created_by=reception)
    invoices.generate_for_contract(c, through=date(2026, 9, 28))
    return c


def test_extend_schedule_rules(partial_contract, reception):
    c = partial_contract
    created = invoices.extend_schedule(c, start_month=date(2026, 10, 1), months=2, due_day=5, by=reception)
    assert [(i.month.month, i.due_date.day, i.amount) for i in created] == [
        (10, 5, D("1750000")), (11, 5, D("1750000"))]
    with pytest.raises(ValidationError):  # bo'shliq: yanvar, lekin dekabr yo'q
        invoices.extend_schedule(c, start_month=date(2027, 1, 1), months=1, by=reception)
    with pytest.raises(ValidationError):  # shartnoma muddatidan tashqari
        invoices.extend_schedule(c, start_month=date(2026, 12, 1), months=12, by=reception)
    with pytest.raises(ValidationError):  # allaqachon bor
        invoices.extend_schedule(c, start_month=date(2026, 10, 1), months=1, by=reception)


def test_advance_payment_needs_created_months(partial_contract, student, reception):
    assert payments.outstanding(student) == D("933000")
    with pytest.raises(ValidationError):
        pay(student, "1000000", reception)  # oktabr hali yaratilmagan
    invoices.extend_schedule(partial_contract, start_month=date(2026, 10, 1), months=1, by=reception)
    pay(student, "2683000", reception)
    assert payments.outstanding(student) == 0


def test_monthly_command_is_idempotent(partial_contract, student):
    assert invoices.generate_due_months(date(2026, 11, 1)) == 2  # okt + noy
    assert invoices.generate_due_months(date(2026, 11, 2)) == 0
    student.status = Student.Status.LEFT
    student.save()
    assert invoices.generate_due_months(date(2026, 12, 1)) == 0  # ketgan o'quvchiga yangi oy yo'q


def test_schedule_dialog_view(staff_client, partial_contract, student):
    detail = staff_client.get(reverse("people:student_detail", args=[student.pk]), {"view": "payment"})
    form = detail.context["schedule_form"]
    assert form.initial["start_month"] == date(2026, 10, 1) and len(form.fields["months"].choices) == 9
    resp = staff_client.post(reverse("finance:schedule_create", args=[student.pk]),
                             {"start_month": "2026-10", "months": "3", "due_date": ""})
    assert resp.url.endswith("?view=payment#grafik")
    assert student.invoices.count() == 4


def test_first_month_due_date_gives_grace_after_late_start(partial_contract):
    sep = partial_contract.invoices.get(month=date(2026, 9, 1))
    assert sep.due_date == date(2026, 9, 25)  # 15-sentabrda kelgan: 10-sana emas, +10 kun


def test_receipt_page_and_guardian_payments_tab(staff_client, contract, student, reception):
    p = pay(student, "933000", reception, method=Payment.Method.TRANSFER)
    resp = staff_client.get(reverse("finance:payment_receipt", args=[p.pk]), {"next": "https://evil.example/"})
    html = resp.content.decode()
    assert resp.status_code == 200 and p.number in html and money(933000) in html
    assert "evil.example" not in resp.context["back_url"]  # ?next= faqat o'z saytimiz
    guardian = student.guardians.first()
    tab = staff_client.get(reverse("people:guardian_detail", args=[guardian.pk]), {"tab": "payments"})
    assert tab.context["payments"][0].payment_trx[0].kind == Transaction.Kind.PAYMENT
    debts = staff_client.get(reverse("people:guardian_detail", args=[guardian.pk]), {"tab": "debts"})
    assert len(debts.context["invoices"]) == 10 and debts.context["debt_count"] == 9  # to'langani ham ko'rinadi


def test_cashbox_overview_numbers(staff_client, contract, student, reception):
    session = cash.open_session(student.branch, by=reception)
    pay(student, "500000", reception, method=Payment.Method.CASH)
    pay(student, "433000", reception)  # karta: 0.2% komissiya → net 432 134
    o = selectors.cashbox_overview(student.branch, timezone.localdate())
    assert o["cash_balance"] == D("500000") and o["total_balance"] == D("500000") + o["bank_balance"]
    assert o["today"]["income_count"] == 2 and o["session"] == session and o["stable"]
    kinds = {k["code"]: k for k in o["by_kind"]}
    assert kinds["cash"]["income"] == D("500000") and kinds["card"]["net"] == D("433000") - D("866")
    assert "bank" not in kinds  # harakati yo'q hisob ko'rsatilmaydi
    resp = staff_client.get(reverse("finance:cashbox"), {"date": "2020-01-01"})
    assert resp.context["o"]["today"]["income_count"] == 0 and not resp.context["is_today"]


def test_payment_list_filters_stats_and_export(staff_client, contract, student, reception):
    from io import BytesIO

    from openpyxl import load_workbook
    p1 = pay(student, "933000", reception, method=Payment.Method.TRANSFER)
    p2 = pay(student, "100000", reception, method=Payment.Method.TRANSFER)
    payments.reverse_payment(p2, reason="Xato", by=reception)
    url = reverse("finance:payment_list")
    resp = staff_client.get(url)  # sukut: bugungi
    s = resp.context["stats"]
    assert (s["total"], s["count"], s["storno_count"]) == (D("933000"), 1, 1)
    assert "date_from=" in resp.context["querystring"]
    assert staff_client.get(url, {"method": "cash"}).context["page"].paginator.count == 0
    assert staff_client.get(url, {"q": p1.number}).context["page"].paginator.count == 1
    assert staff_client.get(url, {"date_from": "2020-01-01", "date_to": "2020-01-02"}).context["stats"]["count"] == 0
    rows = list(load_workbook(BytesIO(staff_client.get(reverse("finance:payment_export")).content))
                .active.iter_rows(min_row=2, values_only=True))
    assert sorted(r[10] for r in rows) == ["Muvaffaqiyatli", "Storno"]


def test_payment_list_reception_only(client, branch, teacher_user):
    client.force_login(teacher_user)
    assert client.get(reverse("finance:payment_list")).status_code in (302, 403)


def test_transaction_list_filters_and_stats(staff_client, contract, student, reception):
    from io import BytesIO

    from openpyxl import load_workbook
    pay(student, "433000", reception)  # karta: komissiya 866
    p2 = pay(student, "100000", reception, method=Payment.Method.TRANSFER)
    payments.reverse_payment(p2, reason="Xato", by=reception)
    url = reverse("finance:transaction_list")
    s = staff_client.get(url).context["stats"]  # sukut: bugungi
    assert (s["count"], s["income"], s["outcome"], s["commission"]) == (3, D("533000"), D("100000"), D("866"))
    assert s["net"] == D("533000") - D("866") - D("100000")
    assert {k["kind"] for k in s["by_kind"]} == {"terminal", "bank"}
    assert staff_client.get(url, {"status": "cancelled"}).context["stats"]["count"] == 1  # storno qilingan to'lov
    assert staff_client.get(url, {"status": "storno"}).context["stats"]["count"] == 1  # teskari yozuv
    assert staff_client.get(url, {"account_kind": ["terminal", "bank"]}).context["stats"]["count"] == 3
    assert staff_client.get(url, {"category": ["tuition"]}).context["stats"]["count"] == 2
    assert staff_client.get(url, {"category": ["storno", "expense"]}).context["stats"]["count"] == 1
    assert staff_client.get(url, {"card_network": "humo"}).context["stats"]["count"] == 1
    assert staff_client.get(url, {"min_amount": "200000"}).context["stats"]["count"] == 1
    trx = Transaction.objects.filter(kind="payment").first()
    assert staff_client.get(url, {"q": f"TRX-{trx.pk}"}).context["stats"]["count"] == 1
    rows = list(load_workbook(BytesIO(staff_client.get(reverse("finance:transaction_export")).content))
                .active.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 3 and sorted(r[14] for r in rows) == ["Bekor qilingan", "Muvaffaqiyatli", "Storno"]


def test_bank_account_overview(staff_client, contract, student, reception):
    pay(student, "433000", reception)  # terminal: komissiya 866
    pay(student, "500000", reception, method=Payment.Method.TRANSFER)
    resp = staff_client.get(reverse("finance:bank_account"))  # sukut: oy boshidan bugungacha
    o = resp.context["o"]
    assert o["bank"]["net"] == D("500000") and o["terminal"]["commission"] == D("866")
    assert o["terminal"]["net"] == D("433000") - D("866") and o["total"] == D("933000") - D("866")
    assert o["all_time"] == o["total"] and resp.context["page"].paginator.count == 2
    empty = staff_client.get(reverse("finance:bank_account"), {"date_from": "2020-01-01", "date_to": "2020-01-31"})
    assert empty.context["o"]["total"] == 0 and empty.context["o"]["all_time"] == o["all_time"]


def test_invoice_list_stats_and_filters(staff_client, contract, student, reception):
    pay(student, "1000000", reception)  # sen to'liq (933 000) + okt qisman (67 000)
    url = reverse("finance:invoice_list")
    s = staff_client.get(url).context["stats"]
    assert s["count"] == 10 and s["paid"] == D("1000000") and s["waived"] == D("817000")
    assert s["total"] - s["waived"] - s["paid"] == s["remaining"]  # asl tizimdagi formula
    assert staff_client.get(url, {"month": "2026-10"}).context["page"].paginator.count == 1
    assert staff_client.get(url, {"state": "paid"}).context["page"].paginator.count == 1
    resp = staff_client.get(url, {"state": "overdue", "student_status": "active"})
    assert resp.context["active_filters"] == {"state", "student_status"}
    assert staff_client.get(reverse("finance:invoice_export"), {"month": "2026-09"}).status_code == 200



def test_invoice_list_waived_cancelled_and_student_status(staff_client, contract, student, reception):
    url = reverse("finance:invoice_list")
    assert staff_client.get(url, {"state": "waived"}).context["page"].paginator.count == 1  # 15-sentabrda kelgan
    contract.invoices.filter(month=date(2027, 6, 1)).update(status=Invoice.Status.CANCELLED)
    assert staff_client.get(url).context["stats"]["count"] == 9  # bekor qilingan statistikaga kirmaydi
    assert staff_client.get(url, {"state": "cancelled"}).context["page"].paginator.count == 1
    labels = [label for _, label in staff_client.get(url).context["form"].fields["student_status"].choices]
    assert labels == ["O'quvchi holati", "Faol", "Nofaol", "Bitirgan", "O'qishdan chiqarilgan", "Ta'tilda", "Ketgan"]


# ---------- Qarzdorlar ----------

def test_debtors_list_panels_and_pay_flow(staff_client, partial_contract, student, reception, monkeypatch):
    monkeypatch.setattr(timezone, "localdate", lambda *a: date(2026, 9, 28))  # "joriy oy" — sentabr
    url = reverse("finance:debtor_list")
    invoices.extend_schedule(partial_contract, start_month=date(2026, 10, 1), months=3, by=reception)
    resp = staff_client.get(url)
    row = resp.context["page"].object_list[0]
    # faqat kelgan oylar (joriy — sentabr) qarz; oldindan yaratilgan okt–dek hisobga kirmaydi
    assert (row.debt_months, row.remaining) == (1, D("933000")) and resp.context["stats"]["students"] == 1
    info = staff_client.get(reverse("finance:debtor_panel", args=[student.pk, "info"]))
    assert info.context["remaining"] == D("933000")
    pay_panel = staff_client.get(reverse("finance:debtor_panel", args=[student.pk, "pay"]), {"next": url})
    assert pay_panel.context["next"] == url and pay_panel.context["form"].initial["amount"] == 933000
    resp = staff_client.post(reverse("finance:payment_create", args=[student.pk]), {
        "amount": "933000", "method": "transfer", "next": url, "print": "1"})
    p = Payment.objects.get()
    assert resp.url.startswith(reverse("finance:payment_receipt", args=[p.pk]))  # kvitansiyaga
    assert staff_client.get(url).context["page"].paginator.count == 0  # endi qarzdor emas


def test_payment_date_rules(partial_contract, student, reception, monkeypatch):
    from datetime import datetime, time
    monkeypatch.setattr(timezone, "localdate", lambda *a: date(2026, 9, 28))
    at = lambda d: timezone.make_aware(datetime.combine(d, time(12)))  # noqa: E731
    with pytest.raises(ValidationError):  # kelajak
        payments.accept_payment(student=student, amount=D("1000"), method="transfer", by=reception,
                                paid_at=at(date(2026, 9, 30)))
    cash.open_session(student.branch, by=reception)
    with pytest.raises(ValidationError):  # naqd — faqat bugun
        payments.accept_payment(student=student, amount=D("1000"), method="cash", by=reception,
                                paid_at=at(date(2026, 9, 20)))
    p = payments.accept_payment(student=student, amount=D("1000"), method="transfer", by=reception,
                                paid_at=at(date(2026, 9, 20)))
    assert timezone.localtime(p.paid_at).date() == date(2026, 9, 20)


def test_budget_page_closed_for_now(staff_client, settings):
    settings.BUDGET_PAGE_ENABLED = False
    resp = staff_client.get(reverse("finance:budget_edit"))
    assert resp.url == reverse("core:section", args=["budget"])
    page = staff_client.get(resp.url)
    assert "Bu bo'lim tayyorlanmoqda" in page.content.decode()
