"""Excel (.xlsx) eksport: sarlavha qatori, ustun kengliklari, formula-injection himoyasi."""

from collections.abc import Iterable, Sequence
from io import BytesIO

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_WIDTH = 48


def _cell_value(value):
    # "=" bilan boshlangan matnni openpyxl formula deb yozadi — foydalanuvchi ma'lumoti hech qachon formula bo'lmasin.
    if isinstance(value, str) and value.startswith("="):
        return "'" + value
    return value


def build_workbook(title: str, headers: Sequence[str], rows: Iterable[Sequence]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append(list(headers))
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E8EEF0")
    widths = [len(h) for h in headers]
    for row in rows:
        values = [_cell_value(v) for v in row]
        ws.append(values)
        for i, v in enumerate(values):
            widths[i] = max(widths[i], len(str(v)) if v is not None else 0)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(w + 2, MAX_WIDTH)
    ws.freeze_panes = "A2"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def xlsx_response(filename: str, content: bytes) -> HttpResponse:
    response = HttpResponse(content, content_type=XLSX_TYPE)
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
