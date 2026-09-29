from datetime import date
from io import BytesIO

import pytest
from django.urls import reverse
from openpyxl import load_workbook

from apps.accounts.models import Role
from apps.contracts.models import Contract
from apps.people.models import Guardian, Student, StudentGuardian
from conftest import make_staff

FORM = {"full_tariff": "1850000", "discount_percent": "0", "start_date": "2026-09-02", "end_date": "2027-06-30"}


def make(branch, school_class=None, **kw):
    data = {"last_name": "Karimov", "first_name": "Ali", "birth_date": date(2016, 5, 1), "gender": "M",
            "grade": 4, "joined_at": date(2026, 9, 2), **kw}
    return Student.objects.create(branch=branch, school_class=school_class, **data)


@pytest.fixture
def students(branch, school_class):
    return {
        "a": make(branch, school_class, first_name="Ali", in_erp=True),
        "b": make(branch, school_class, first_name="Bobur", gender="F", joined_at=date(2026, 8, 1)),
        "c": make(branch, first_name="Doniyor", grade=7, status=Student.Status.LEFT),
    }


def with_contract(client, student):
    guardian = Guardian.objects.create(last_name="Karimov", first_name="Anvar", phone="998901112233")
    StudentGuardian.objects.create(student=student, guardian=guardian, relation="father", is_primary=True)
    client.post(reverse("contracts:contract_create", args=[student.pk]), FORM)
    assert Contract.objects.filter(student=student).exists()


def names(resp):
    return {s.first_name for s in resp.context["page"].object_list}


def test_default_list_shows_only_active(staff_client, students):
    resp = staff_client.get(reverse("people:student_list"))
    assert names(resp) == {"Ali", "Bobur"}
    assert resp.context["active_filters"] == {"status"}
    assert not resp.context["has_filters"]


@pytest.mark.parametrize(("query", "expected"), [
    ({"in_erp": "1"}, {"Ali"}),
    ({"in_erp": "0"}, {"Bobur"}),
    ({"gender": "F"}, {"Bobur"}),
    ({"joined_from": "2026-09-01"}, {"Ali"}),
    ({"status": "", "grade": "7"}, {"Doniyor"}),
    ({"status": ""}, {"Ali", "Bobur", "Doniyor"}),
])
def test_filters(staff_client, students, query, expected):
    resp = staff_client.get(reverse("people:student_list"), query)
    assert names(resp) == expected
    assert resp.context["has_filters"]


def test_academic_year_filter(staff_client, students, year):
    resp = staff_client.get(reverse("people:student_list"), {"status": "", "academic_year": year.pk})
    assert names(resp) == {"Ali", "Bobur"}  # Doniyor sinfsiz


def test_tariff_filter_uses_contract(staff_client, students):
    with_contract(staff_client, students["a"])
    resp = staff_client.get(reverse("people:student_list"), {"tariff": "1850000.00"})
    assert names(resp) == {"Ali"}


def test_export_respects_filters_and_escapes_formulas(staff_client, branch, students):
    make(branch, first_name="=HYPERLINK(\"x\")")
    resp = staff_client.get(reverse("people:student_export"), {"gender": "M"})
    assert resp.status_code == 200
    assert resp["Content-Disposition"].startswith('attachment; filename="oquvchilar-')
    ws = load_workbook(BytesIO(resp.content)).active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert {r[3] for r in rows} == {"Ali", "'=HYPERLINK(\"x\")"}
    assert ws["A1"].value == "#"


def test_delete_without_history_removes_student(staff_client, students):
    s = students["b"]
    resp = staff_client.post(reverse("people:student_delete", args=[s.pk]), {"qs": "status=active"})
    assert resp.status_code == 302 and resp.url.endswith("?status=active")
    assert not Student.objects.filter(pk=s.pk).exists()


def test_delete_with_contract_only_archives(staff_client, students):
    s = students["a"]
    with_contract(staff_client, s)
    staff_client.post(reverse("people:student_delete", args=[s.pk]))
    s.refresh_from_db()
    assert s.status == Student.Status.LEFT and s.left_at == date.today()


def test_delete_requires_post_and_own_branch(staff_client, other_branch):
    foreign = make(other_branch)
    assert staff_client.get(reverse("people:student_delete", args=[foreign.pk])).status_code == 405
    assert staff_client.post(reverse("people:student_delete", args=[foreign.pk])).status_code == 404
    assert Student.objects.filter(pk=foreign.pk).exists()


def test_teacher_cannot_export_or_delete(client, teacher_user, students):
    client.force_login(teacher_user)
    assert client.get(reverse("people:student_export")).status_code in (302, 403)
    client.post(reverse("people:student_delete", args=[students["b"].pk]))
    assert Student.objects.filter(pk=students["b"].pk).exists()


# ---------- O'quvchi kartasi ----------

def test_detail_academic_shows_siblings_and_masks_documents(staff_client, branch, students):
    a, b = students["a"], students["b"]
    guardian = Guardian.objects.create(last_name="Karimov", first_name="Anvar", phone="998901112233")
    for s in (a, b):
        StudentGuardian.objects.create(student=s, guardian=guardian, relation="father", is_primary=True)
    a.birth_certificate = "i-fr 1055437"
    a.region = "Farg'ona viloyati"
    a.save()
    a.refresh_from_db()
    assert a.birth_certificate == "I-FR1055437"
    resp = staff_client.get(reverse("people:student_detail", args=[a.pk]))
    html = resp.content.decode()
    assert resp.context["view"] == "academic"
    assert [s.pk for s in resp.context["siblings"]] == [b.pk]
    assert "I-FR1055437" not in html and "5437" in html  # niqoblangan
    assert "Farg&#x27;ona viloyati" in html


def test_detail_payment_view_only_for_reception(client, staff_client, branch, students):
    a = students["a"]
    resp = staff_client.get(reverse("people:student_detail", args=[a.pk]), {"view": "payment"})
    assert resp.context["view"] == "payment" and "charged" in resp.context
    head = make_staff(branch, Role.HEAD_TEACHER, "998900000109")
    client.force_login(head)
    resp = client.get(reverse("people:student_detail", args=[a.pk]), {"view": "payment"})
    assert resp.context["view"] == "academic" and "charged" not in resp.context
    assert client.get(reverse("finance:student_invoices_export", args=[a.pk])).status_code in (302, 403)


def test_student_form_normalizes_documents(staff_client, students):
    s = students["a"]
    data = {"last_name": s.last_name, "first_name": s.first_name, "birth_date": "2016-05-01", "gender": "M",
            "grade": "4", "school_class": s.school_class_id, "status": "active", "joined_at": "2026-09-02",
            "passport": "ad 1234567", "birth_certificate": "i-fr-1055437", "district": "Oltiariq tumani"}
    resp = staff_client.post(reverse("people:student_update", args=[s.pk]), data)
    assert resp.status_code == 302, resp.content.decode()[:800]
    s.refresh_from_db()
    assert (s.passport, s.birth_certificate, s.district) == ("AD1234567", "I-FR-1055437", "Oltiariq tumani")
    assert s.passport_index  # blind index yozildi


# ---------- Tahrirlash sahifasi (o'quvchi + vasiylar bitta formada) ----------

def edit_payload(student, **extra):
    data = {"last_name": student.last_name, "first_name": student.first_name, "birth_date": "2016-05-01",
            "gender": "M", "grade": "4", "school_class": student.school_class_id or "", "status": "active",
            "joined_at": "2026-09-02", "region": "Farg'ona viloyati"}
    for link in student.guardian_links.select_related("guardian"):
        g = link.guardian
        data.update({f"g{link.pk}-last_name": g.last_name, f"g{link.pk}-first_name": g.first_name,
                     f"g{link.pk}-phone": g.phone[3:]})
    return {**data, **extra}


@pytest.fixture
def family(branch, students):
    a, b = students["a"], students["b"]
    mother = Guardian.objects.create(last_name="Karimova", first_name="Nodira", phone="998901112233")
    father = Guardian.objects.create(last_name="Karimov", first_name="Anvar", phone="998901114455")
    primary = StudentGuardian.objects.create(student=a, guardian=mother, relation="mother", is_primary=True)
    StudentGuardian.objects.create(student=b, guardian=mother, relation="mother", is_primary=True)
    second = StudentGuardian.objects.create(student=a, guardian=father, relation="father")
    return a, primary, second


def test_edit_page_saves_student_and_guardians_together(staff_client, family):
    a, primary, _ = family
    resp = staff_client.get(reverse("people:student_update", args=[a.pk]))
    link = next(link for link, _ in resp.context["guardian_forms"] if link.pk == primary.pk)
    assert link.children_total == 2  # "2 o'quvchi" ogohlantirishi
    resp = staff_client.post(reverse("people:student_update", args=[a.pk]), edit_payload(
        a, district="Oltiariq tumani",
        **{f"g{primary.pk}-position": "O'qituvchi", f"g{primary.pk}-first_name": "Dilnoza"}))
    assert resp.status_code == 302, resp.content.decode()[:600]
    a.refresh_from_db()
    primary.guardian.refresh_from_db()
    assert a.district == "Oltiariq tumani" and a.region == "Farg'ona viloyati"
    assert (primary.guardian.first_name, primary.guardian.position) == ("Dilnoza", "O'qituvchi")


def test_edit_invalid_guardian_rolls_back_everything(staff_client, family):
    a, primary, _ = family
    resp = staff_client.post(reverse("people:student_update", args=[a.pk]), edit_payload(
        a, first_name="Yangi", **{f"g{primary.pk}-pinfl": "123"}))
    assert resp.status_code == 200
    a.refresh_from_db()
    assert a.first_name == "Ali"


def test_edit_adds_new_guardian_with_dedupe(staff_client, branch, family):
    a, _, _ = family
    payload = edit_payload(a, **{"new-enabled": "1", "new-r-relation": "other", "new-last_name": "Karimov",
                                 "new-first_name": "Bobur", "new-phone": "90 777 66 55"})
    assert staff_client.post(reverse("people:student_update", args=[a.pk]), payload).status_code == 302
    assert a.guardian_links.filter(relation="other", guardian__first_name="Bobur").exists()
    # Blok ochilmagan bo'lsa (new-enabled=0) bo'sh vasiy maydonlari e'tiborga olinmaydi
    assert staff_client.post(reverse("people:student_update", args=[a.pk]),
                             edit_payload(a, **{"new-enabled": "0"})).status_code == 302


def test_detach_guardian_rules(staff_client, family):
    a, primary, second = family
    url = lambda link: reverse("people:student_detach_guardian", args=[a.pk, link.pk])  # noqa: E731
    staff_client.post(url(primary))
    assert StudentGuardian.objects.filter(pk=primary.pk).exists()  # asosiy aloqa ajratilmaydi
    resp = staff_client.post(url(second))
    assert resp.url.endswith("/edit/#vasiylar")
    assert not StudentGuardian.objects.filter(pk=second.pk).exists()
    assert Guardian.objects.filter(pk=second.guardian_id).exists()  # vasiy yozuvi qoladi


# ---------- Ota-onalar: ro'yxat (faqat ko'rish), karta, Excel ----------

def test_guardian_list_is_view_only_and_filters(staff_client, family):
    resp = staff_client.get(reverse("people:guardian_list"))
    html = resp.content.decode()
    assert "/edit/" not in html and resp.context["page"].paginator.count == 2
    resp = staff_client.get(reverse("people:guardian_list"), {"relation": "father"})
    assert [g.first_name for g in resp.context["page"].object_list] == ["Anvar"]
    assert resp.context["has_filters"] and resp.context["active_filters"] == {"relation"}


def test_guardian_edit_page_removed():
    from django.urls import NoReverseMatch
    with pytest.raises(NoReverseMatch):
        reverse("people:guardian_update", args=[1])


def test_guardian_detail_tabs_and_finance_visibility(client, staff_client, branch, family):
    a, primary, _ = family
    url = reverse("people:guardian_detail", args=[primary.guardian_id])
    resp = staff_client.get(url, {"tab": "payments"})
    assert resp.context["tab"] == "payments" and resp.context["total_debt"] == 0
    assert staff_client.get(url, {"tab": "debts"}).status_code == 200
    assert len(resp.context["links"]) == 2 and resp.context["is_primary"]
    client.force_login(make_staff(branch, Role.HEAD_TEACHER, "998900000112"))
    resp = client.get(url, {"tab": "payments"})
    assert resp.context["tab"] == "children" and "total_debt" not in resp.context  # moliya faqat Reception


def test_guardian_export_hides_documents(staff_client, family):
    _, primary, _ = family
    g = primary.guardian
    g.pinfl = "31234567890123"
    g.save()
    resp = staff_client.get(reverse("people:guardian_export"))
    ws = load_workbook(BytesIO(resp.content)).active
    values = [c for row in ws.iter_rows(values_only=True) for c in row]
    assert "Karimova" in values and "31234567890123" not in values


def test_extended_relations_count_as_other_guardians(staff_client, branch, students):
    g = Guardian.objects.create(last_name="Karimov", first_name="Buva", phone="998901119999")
    StudentGuardian.objects.create(student=students["a"], guardian=g, relation="grandpa", is_primary=True)
    resp = staff_client.get(reverse("people:guardian_list"), {"relation": "grandpa"})
    assert [x.first_name for x in resp.context["page"].object_list] == ["Buva"]
    assert resp.context["stats"]["others"] == 1
    labels = [label for _, label in resp.context["filter_form"].fields["relation"].choices]
    assert labels == ["Munosabat", "Ota", "Ona", "Buva", "Buvi", "Tog'a", "Xola", "Boshqa"]
