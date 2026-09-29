import re
from datetime import date
from decimal import Decimal

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse

from apps.accounts.infrastructure.sms import LocmemSMSBackend
from apps.contracts.models import Contract
from apps.people.models import Guardian, Student, StudentGuardian

FORM = {"full_tariff": "1750000", "discount_percent": "0", "start_date": "2026-09-02", "end_date": "2027-06-30"}


@pytest.fixture(autouse=True)
def clear_outbox():
    LocmemSMSBackend.outbox.clear()


@pytest.fixture
def guardian(db):
    return Guardian.objects.create(last_name="Aliyev", first_name="Anvar", phone="998901112233")


def make_student(branch, school_class, guardian, first_name="Vali"):
    s = Student.objects.create(branch=branch, last_name="Aliyev", first_name=first_name, birth_date=date(2016, 5, 1),
                               gender="M", grade=4, school_class=school_class, joined_at=date(2026, 9, 2))
    StudentGuardian.objects.create(student=s, guardian=guardian, relation="father", is_primary=True)
    return s


@pytest.fixture
def student(branch, school_class, guardian):
    return make_student(branch, school_class, guardian)


def create(client, student, **data):
    return client.post(reverse("contracts:contract_create", args=[student.pk]), {**FORM, **data})


def sent_code():
    return re.search(r"kodi: (\d{6})", LocmemSMSBackend.outbox[-1][1]).group(1)


@pytest.fixture
def contract(staff_client, student):
    create(staff_client, student)
    return Contract.objects.get()


def test_create_prefills_class_tariff(staff_client, student):
    page = staff_client.get(reverse("contracts:contract_create", args=[student.pk])).content.decode()
    assert 'value="1750000' in page


def test_monthly_fee_with_discount(staff_client, student):
    create(staff_client, student, discount_percent="10", discount_reason="Ikkinchi farzand")
    c = Contract.objects.get()
    assert c.monthly_fee == Decimal("1575000")
    assert c.number == "CTR-2026-001" and c.status == Contract.Status.DRAFT
    assert c.months == 10 and c.class_name == "4-A"


def test_discount_requires_reason(staff_client, student):
    resp = create(staff_client, student, discount_percent="10")
    assert resp.status_code == 200 and not Contract.objects.exists()


def test_dates_must_be_inside_academic_year(staff_client, student):
    create(staff_client, student, start_date="2025-09-01")
    assert not Contract.objects.exists()


def test_one_active_contract_per_student_year(staff_client, student, contract):
    resp = create(staff_client, student)
    assert resp.url == reverse("contracts:contract_detail", args=[contract.pk])
    assert Contract.objects.count() == 1


def test_full_sms_confirmation_flow(staff_client, contract):
    staff_client.post(reverse("contracts:contract_send_code", args=[contract.pk]))
    contract.refresh_from_db()
    assert contract.status == Contract.Status.SENT
    phone, text = LocmemSMSBackend.outbox[-1]
    assert phone == "998901112233" and contract.number in text and "1 750 000" in text

    staff_client.post(reverse("contracts:contract_confirm", args=[contract.pk]), {"code": "000000"})
    contract.refresh_from_db()
    assert contract.status == Contract.Status.SENT

    staff_client.post(reverse("contracts:contract_confirm", args=[contract.pk]), {"code": sent_code()})
    contract.refresh_from_db()
    assert contract.status == Contract.Status.SIGNED
    assert contract.signed_phone == "998901112233" and contract.signed_at


def test_code_is_bound_to_its_contract(staff_client, branch, school_class, guardian, contract):
    """Bitta ota-onaning ikki farzandi: A shartnoma kodi B shartnomani imzolay olmaydi."""
    sibling = make_student(branch, school_class, guardian, first_name="Madina")
    create(staff_client, sibling)
    other = Contract.objects.exclude(pk=contract.pk).get()

    staff_client.post(reverse("contracts:contract_send_code", args=[contract.pk]))
    code_a = sent_code()
    staff_client.post(reverse("contracts:contract_send_code", args=[other.pk]))
    staff_client.post(reverse("contracts:contract_confirm", args=[other.pk]), {"code": code_a})
    other.refresh_from_db()
    assert other.status == Contract.Status.SENT  # A kodi B uchun ishlamadi


def test_signed_contract_not_editable(staff_client, contract):
    Contract.objects.filter(pk=contract.pk).update(status=Contract.Status.SIGNED)
    staff_client.post(reverse("contracts:contract_update", args=[contract.pk]), {**FORM, "full_tariff": "1"})
    contract.refresh_from_db()
    assert contract.full_tariff == Decimal("1750000")


def test_cancel_requires_reason_then_allows_new_contract(staff_client, student, contract):
    staff_client.post(reverse("contracts:contract_cancel", args=[contract.pk]), {"reason": ""})
    contract.refresh_from_db()
    assert contract.status == Contract.Status.DRAFT
    staff_client.post(reverse("contracts:contract_cancel", args=[contract.pk]), {"reason": "Boshqa maktabga ketdi"})
    contract.refresh_from_db()
    assert contract.status == Contract.Status.CANCELLED
    create(staff_client, student)
    assert Contract.objects.exclude(status=Contract.Status.CANCELLED).count() == 1


def test_scan_must_be_real_pdf_and_is_private(staff_client, contract, settings):
    fake = SimpleUploadedFile("x.pdf", b"<script>alert(1)</script>", content_type="application/pdf")
    staff_client.post(reverse("contracts:contract_upload_scan", args=[contract.pk]), {"scan": fake})
    contract.refresh_from_db()
    assert not contract.scan

    pdf = SimpleUploadedFile("../../etc/passwd.pdf", b"%PDF-1.4 test", content_type="application/pdf")
    staff_client.post(reverse("contracts:contract_upload_scan", args=[contract.pk]), {"scan": pdf})
    contract.refresh_from_db()
    assert contract.scan.name == f"contracts/{contract.academic_year_id}/{contract.number}.pdf"
    assert str(settings.PRIVATE_MEDIA_ROOT) in contract.scan.path  # ommaviy /media/ emas

    resp = staff_client.get(reverse("contracts:contract_scan", args=[contract.pk]))
    assert resp.status_code == 200 and resp["Content-Type"] == "application/pdf"
    assert "no-store" in resp["Cache-Control"]


def test_teacher_cannot_access_contracts(client, teacher_user, contract):
    client.force_login(teacher_user)
    assert client.get(reverse("contracts:contract_list")).status_code == 403
    assert client.get(reverse("contracts:contract_scan", args=[contract.pk])).status_code == 403


def test_contract_sms_limited_per_guardian(staff_client, contract, settings):
    settings.OTP_RESEND_SECONDS = 0
    for _ in range(6):
        staff_client.post(reverse("contracts:contract_send_code", args=[contract.pk]))
    assert len(LocmemSMSBackend.outbox) == 3  # soatiga 3 ta


def test_list_stats(staff_client, contract):
    page = staff_client.get(reverse("contracts:contract_list")).content.decode()
    assert contract.number in page


def test_list_filters_year_class_and_expired(staff_client, contract, year, school_class, monkeypatch):
    url = reverse("contracts:contract_list")
    ids = lambda resp: [c.pk for c in resp.context["page"].object_list]  # noqa: E731
    assert ids(staff_client.get(url, {"academic_year": year.pk, "school_class": school_class.pk})) == [contract.pk]
    resp = staff_client.get(url, {"status": "expired"})
    assert ids(resp) == [] and resp.context["active_filters"] == {"status"}
    Contract.objects.filter(pk=contract.pk).update(status=Contract.Status.SIGNED)
    from django.utils import timezone
    monkeypatch.setattr(timezone, "localdate", lambda *a: date(2027, 7, 1))  # shartnoma 30.06.2027 da tugagan
    assert ids(staff_client.get(url, {"status": "expired"})) == [contract.pk]
    labels = [f["label"] for f in resp.context["filter_menu"]]
    assert labels == ["Barcha yillar", "Barcha sinflar", "Barcha holatlar"]


def test_export_respects_filters(staff_client, contract, year):
    from io import BytesIO

    from openpyxl import load_workbook
    resp = staff_client.get(reverse("contracts:contract_export"), {"academic_year": year.pk})
    assert resp["Content-Disposition"].startswith('attachment; filename="shartnomalar-')
    rows = list(load_workbook(BytesIO(resp.content)).active.iter_rows(min_row=2, values_only=True))
    assert [r[1] for r in rows] == [contract.number] and rows[0][8] == 1750000
    resp = staff_client.get(reverse("contracts:contract_export"), {"status": "cancelled"})
    assert load_workbook(BytesIO(resp.content)).active.max_row == 1  # faqat sarlavha
